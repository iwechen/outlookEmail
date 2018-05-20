# coding:utf-8
'''
Created on 2018年5月15日
@author: chenwei
Email:iwechen123@gmail.com
'''
import requests
import pymongo
import sys
import os
import csv
import xlwt
from docx import Document
from docx.shared import Pt
from docx.shared import Inches
from docx.oxml.ns import qn
import json
import re
import time,datetime
import threading


class MongoToExcel(object):
    def __init__(self):
        self.client = pymongo.MongoClient(host='127.0.0.1', port=27017)
        self.db = self.client['outlook']
        self.collection1 = self.db['data']

        self.iamge_list = list()

    def save_to_excel(self,end_time, data_li,data_li_str):
        end_time = end_time.split(' ')[0]
        path = sys.path[0]+'/%s/'%end_time
        folder = os.path.exists(path)
        if not folder:
            os.makedirs(path)
        # 创建一个Workbook对象，这就相当于创建了一个Excel文件
        book = xlwt.Workbook(encoding='utf-8', style_compression=0)
        # 1.订单概括表
        sheet = book.add_sheet('订单概括', cell_overwrite_ok=True)
        order_li = list()
        order_li_sign = list()
        data_li_order = data_li
        for data_dic in data_li_order:
            sign = data_dic.pop('sign')
            data_dic.pop('image')
            del data_dic['产品编号']
            data_dic.pop('产品金额')
            data_dic.pop('产品重量')
            data_dic.pop('订购数量')
            data_dic.pop('报')
            if sign not in order_li_sign:
                order_li_sign.append(sign)
                order_li.append(data_dic)
        # 写入数据
        for index, data in enumerate(order_li):
            if index == 0:
                for i, k in enumerate(list(data.keys())):
                    sheet.write(index, i, k)
                for i, k in enumerate(list(data.values())):
                    sheet.write(index+1, i, k)
            else:
                for i, k in enumerate(list(data.values())):
                    sheet.write(index+1, i, k)
        # 2.订单详情表
        sheet = book.add_sheet('订单详情', cell_overwrite_ok=True)
        data_li_detail = json.loads(data_li_str)
        order_li_detail = list()
        for data_dic in data_li_detail:
            data_dic.pop('sign')
            data_dic.pop('image')
            data_dic.pop('产品编号集合')
            data_dic.pop('产品重量集合')
            data_dic.pop('订购总数')
            data_dic.pop('实付总额')
            data_dic.pop('报')
            order_li_detail.append(data_dic)
        # 写入数据
        for index, data in enumerate(order_li_detail):
            if index == 0:
                for i, k in enumerate(list(data.keys())):
                    sheet.write(index, i, k)
                for i, k in enumerate(list(data.values())):
                    sheet.write(index+1, i, k)
            else:
                for i, k in enumerate(list(data.values())):
                    sheet.write(index+1, i, k)

        book.save(path+'订单列表.xls')
        print('Excel successful!!!')

    def save_to_word(self,end_time, data_li_str):
        # 打开文档
        data_li = json.loads(data_li_str)
        data_li_dic = list()
        data_li_sign = list()
        for data in data_li:
            sign = data['sign'] 
            if sign not in data_li_sign:
                data_li_sign.append(sign)
                data_li_dic.append(data)
        word_str = ''
        for data in data_li_dic:
            a = data['订单日期']
            b = data['产品编号集合']
            c = data['返邮费金额']
            d = data['报']
            e = data['邮寄地址']
            ww = '%s 订单信息\n %s (%d)报：%d \n %s\r\n' % (a, b, c,d,e)
            word_str += ww
        document = Document()
        paragraph = document.add_paragraph(word_str)
        end_time = end_time.split(' ')[0]
        path = sys.path[0]+'/%s/'%end_time
        document.save(path+'订单列表信息.docx')
        print('Word successful!!!')

    def time_str_to_dtime(self,timeStr):
        ltime = time.localtime(time.mktime(time.strptime(timeStr, "%Y-%m-%d %H:%M")))
        timeStr = time.strftime("%Y-%m-%d %H:%M", ltime)
        date_time = datetime.datetime.strptime(timeStr, '%Y-%m-%d %H:%M')
        return date_time

    def find_mongo(self):
        sta_time = '2018-05-1 11:00'
        sta_temp = self.time_str_to_dtime(sta_time)
        end_time = '2018-05-20 11:00'
        end_temp = self.time_str_to_dtime(end_time)

        data_list = [i for i in self.collection1.find({"date_time":{'$gte':sta_temp,'$lte':end_temp}})]
        data_list_li = list()
        if data_list == []:
            print('当前时间段没有邮件，请启动爬虫获取！！！')
            return
        else:
            for data in data_list:
                data.pop('_id')
                data.pop('date_time')
                data_list_li.append(data)
        data_li_str = json.dumps(data_list_li)
        for data_dict in data_list:
            self.iamge_list.append({"pid": data_dict['产品编号'], 'image': data_dict['image']})

        self.save_to_excel(end_time,data_list,data_li_str)
        
        self.save_to_word(end_time,data_li_str)
        
        self.start_down_image(end_time)

    def start_down_image(self,end_time):
        data_set_pid = list()
        data_set_dic = list()
        for data_dic in self.iamge_list:
            pid = data_dic['pid']
            if pid not in data_set_pid:
                data_set_pid.append(pid)
                data_set_dic.append(data_dic)
        for data_dic in data_set_dic:
            t1 = threading.Thread(target=self.down_load_image,args=(data_dic,end_time))
            t1.start()

    def down_load_image(self,data_dic,end_time):
        ima_name = data_dic['pid']
        url = data_dic['image']
        
        end_time = end_time.split(' ')[0]
        path = sys.path[0]+'/%s/%s/'%(end_time,end_time)
        folder = os.path.exists(path)
        if not folder:
            os.makedirs(path)
        try:
            response = requests.get(url=url).content
        except Exception as e:
            # print(e)
            print(url)
        else:
            print('%s down_load!!!' % ima_name)
            with open(path + ima_name+'.jpg', 'wb') as f:
                f.write(response)

    def run(self):
        self.find_mongo()

    def main(self):
        self.run()

if __name__ == "__main__":
    excel = MongoToExcel()
    excel.main()
